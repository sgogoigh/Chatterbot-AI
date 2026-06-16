"""
services/slide_processor.py — .pptx ingestion + script generation (IMPLEMENTATION.md §14).

Two responsibilities, invoked by two different endpoints:
  * extract()  — POST /api/presentations : parse the .pptx into SlideContent[]
                 (text, notes, embedded images, OCR, full-slide raster).
  * build_scripts() — build : for each (slide, track) ask the LLM for narration,
                 segment into sentences. Pacing budgets come from PacingService.

⚠️ R9: python-pptx extracts embedded pictures but does NOT render a full slide.
We render via LibreOffice headless (pptx -> pdf) then poppler (pdf -> png). The
Dockerfile installs libreoffice-core + poppler-utils for this. Rendering is
best-effort: if the tools are absent, extraction still succeeds without rasters.
"""

from __future__ import annotations

import asyncio
import os
import subprocess

from config import Settings
from models import (
    Persona,
    ScriptSentence,
    SlideContent,
    SlideScript,
    Track,
)
from services.groq_llm_service import GroqLLMService
from utils.logging import get_logger
from utils.text import split_sentences, word_count

log = get_logger(component="slides")


class SlideProcessor:
    """Extracts content from .pptx files and generates per-track narration scripts."""

    def __init__(self, settings: Settings, llm: GroqLLMService):
        """Hold config + the LLM client used for script generation."""
        self.s = settings
        self.llm = llm

    # ----------------------------------------------------------- extraction
    def extract(self, pptx_path: str, job_dir: str) -> list[SlideContent]:
        """Parse a .pptx into a list of :class:`SlideContent` (one per slide).

        Pulls text-frame text, speaker notes, and embedded images; runs OCR over
        the images (tesseract) to recover text baked into diagrams; and attempts a
        full-slide raster for the viewer. Enforces the 1–50 slide bound (§14.1);
        raises ValueError outside that range so the route can return 422.
        """
        from pptx import Presentation

        prs = Presentation(pptx_path)
        slides = list(prs.slides)
        if not (1 <= len(slides) <= self.s.max_slides):
            raise ValueError(f"slide count {len(slides)} outside 1..{self.s.max_slides}")

        images_dir = os.path.join(job_dir, "images")
        os.makedirs(images_dir, exist_ok=True)
        renders = self._render_slides(pptx_path, job_dir)  # may be [] if tools absent

        out: list[SlideContent] = []
        for i, slide in enumerate(slides):
            title = self._title_of(slide)
            body = "\n".join(
                sh.text for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text
            )
            notes = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text
            imgs = self._export_pictures(slide, images_dir, i)
            ocr = self._ocr(imgs)
            out.append(
                SlideContent(
                    index=i, title=title, body_text=body, notes=notes,
                    ocr_text=ocr, image_paths=imgs,
                    render_path=renders[i] if i < len(renders) else None,
                )
            )
        log.info(f"extracted {len(out)} slides")
        return out

    def _title_of(self, slide) -> str | None:
        """Return the slide's title placeholder text, if any."""
        try:
            if slide.shapes.title and slide.shapes.title.text:
                return slide.shapes.title.text
        except Exception:  # noqa: BLE001 - layouts without a title placeholder
            pass
        return None

    def _export_pictures(self, slide, images_dir: str, idx: int) -> list[str]:
        """Save embedded pictures of one slide to disk; return their file paths.

        These feed both the OCR step and (as a fallback) the frontend viewer when
        a full-slide raster isn't available.
        """
        paths: list[str] = []
        for j, shape in enumerate(slide.shapes):
            if getattr(shape, "shape_type", None) == 13:  # MSO_SHAPE_TYPE.PICTURE
                try:
                    blob = shape.image.blob
                    ext = shape.image.ext or "png"
                    p = os.path.join(images_dir, f"slide{idx}_img{j}.{ext}")
                    with open(p, "wb") as f:
                        f.write(blob)
                    paths.append(p)
                except Exception as e:  # noqa: BLE001 - skip unreadable media
                    log.debug(f"image export skipped (slide {idx}): {e}")
        return paths

    def _ocr(self, image_paths: list[str]) -> str:
        """Run tesseract OCR over the given images and join the recovered text.

        Best-effort: OCR failures (missing tesseract, unsupported image) degrade
        to empty text rather than failing extraction.
        """
        if not image_paths:
            return ""
        try:
            import pytesseract
            from PIL import Image
        except Exception:  # noqa: BLE001 - OCR optional
            return ""
        chunks = []
        for p in image_paths:
            try:
                chunks.append(pytesseract.image_to_string(Image.open(p)).strip())
            except Exception:  # noqa: BLE001 - skip bad image
                continue
        return "\n".join(c for c in chunks if c)

    def _render_slides(self, pptx_path: str, job_dir: str) -> list[str]:
        """Render each slide to a PNG via LibreOffice + poppler (R9). Best-effort.

        Returns an ordered list of PNG paths (index-aligned to slides), or an
        empty list if the rendering tools are unavailable — in which case the
        frontend falls back to embedded images / its own chrome.
        """
        try:
            out_dir = os.path.join(job_dir, "renders")
            os.makedirs(out_dir, exist_ok=True)
            # pptx -> pdf
            subprocess.run(
                ["soffice", "--headless", "--convert-to", "pdf", "--outdir", out_dir, pptx_path],
                check=True, capture_output=True, timeout=120,
            )
            pdf = os.path.join(out_dir, os.path.splitext(os.path.basename(pptx_path))[0] + ".pdf")
            # pdf -> png pages
            from pdf2image import convert_from_path

            pages = convert_from_path(pdf, dpi=120)
            paths = []
            for i, page in enumerate(pages):
                pp = os.path.join(out_dir, f"slide{i}.png")
                page.save(pp, "PNG")
                paths.append(pp)
            return paths
        except Exception as e:  # noqa: BLE001 - rendering is optional
            log.warning(f"slide rasterization unavailable: {e}")
            return []

    # ---------------------------------------------------- script generation
    async def build_scripts(
        self,
        slides: list[SlideContent],
        word_budgets: dict[int, int],
        persona: Persona,
        track: Track,
        on_slide_done=None,
    ) -> list[SlideScript]:
        """Generate the narration script for every slide on a single ``track``.

        ``word_budgets`` maps slide_index -> target words (from PacingService).
        For each slide we ask the LLM for narration, retry once if the word count
        is wildly off budget, then segment into :class:`ScriptSentence` objects so
        the PlaybackTracker can map them to audio. ``on_slide_done`` is an optional
        progress callback for the async build status endpoint (§7.1).
        """
        scripts: list[SlideScript] = []
        for slide in slides:
            budget = word_budgets.get(slide.index, 80)
            text = await self.llm.generate_script(slide, budget, persona, track)
            if word_count(text) > budget * 1.5 or word_count(text) < budget * 0.5:
                # One retry: large miss usually means the model ignored the budget.
                text = await self.llm.generate_script(slide, budget, persona, track)
            sentences = [
                ScriptSentence(text=s, word_count=word_count(s)) for s in split_sentences(text)
            ]
            scripts.append(SlideScript(slide_index=slide.index, track=track, sentences=sentences))
            if on_slide_done:
                on_slide_done()
            await asyncio.sleep(0)  # cooperative yield between LLM calls
        return scripts
